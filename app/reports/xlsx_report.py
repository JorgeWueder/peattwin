"""Informe Excel (.xlsx): exportacion cruda de flux_observations + predictions +
tabla completa de metricas por modelo y por fold."""
from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.reports import common as C

_HEADER_FILL = PatternFill("solid", fgColor="ECFDF5")
_HEADER_FONT = Font(bold=True, color="047857")


def _sheet(wb: Workbook, name: str, headers: list[str], rows: list[list]):
    ws = wb.create_sheet(name[:31])
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="left")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    # ancho de columna aproximado
    for c, h in enumerate(headers, start=1):
        maxlen = max([len(str(h))] + [len(str(r[c - 1])) for r in rows[:200] if c - 1 < len(r)] or [0])
        ws.column_dimensions[get_column_letter(c)].width = min(max(maxlen + 2, 10), 40)
    return ws


def _cell(v):
    if isinstance(v, datetime):
        return v.replace(tzinfo=None)
    return v


def build(db: Session, *, site_id: int) -> bytes:
    site = C.get_site(db, site_id)
    wb = Workbook()
    wb.remove(wb.active)  # quita la hoja por defecto

    # ---- info ----
    rows_active = C.model_rows(db)
    winner = next((r for r in rows_active if r.is_active), None)
    _sheet(
        wb, "info",
        ["clave", "valor"],
        [
            ["generado_utc", datetime.now(timezone.utc).replace(tzinfo=None)],
            ["sitio_id", site.id],
            ["sitio_nombre", site.name],
            ["fluxnet_id", site.fluxnet_id or ""],
            ["peatland_type", str(getattr(site.peatland_type, "value", site.peatland_type) or "")],
            ["latitud", site.latitude],
            ["longitud", site.longitude],
            ["modelo_activo", winner.name if winner else ""],
            ["modelo_activo_version", winner.version if winner else ""],
        ],
    )

    # ---- flux_observations ----
    obs = C.flux_rows(db, site_id, None, None)
    _sheet(
        wb, "flux_observations",
        [
            "id", "timestamp", "nee_co2", "fch4", "water_table_depth_cm", "soil_temp_c",
            "soil_water_content", "air_temp_c", "precipitation_mm", "carbon_balance_class",
            "quality_flag",
        ],
        [
            [
                o.id, _cell(o.timestamp), o.nee_co2, o.fch4, o.water_table_depth, o.soil_temp,
                o.soil_water_content, o.air_temp, o.precipitation,
                str(getattr(o.carbon_balance_class, "value", o.carbon_balance_class) or ""),
                o.quality_flag,
            ]
            for o in obs
        ],
    )

    # ---- predictions ----
    preds = C.prediction_rows(db, site_id)
    _sheet(
        wb, "predictions",
        ["id", "escenario", "modelo", "timestamp", "co2_predicho", "ch4_predicho", "clase_predicha"],
        [
            [p["id"], p["escenario"], p["modelo"], _cell(p["timestamp"]),
             p["co2_predicted"], p["ch4_predicted"],
             str(getattr(p["predicted_class"], "value", p["predicted_class"]) or "")]
            for p in preds
        ],
    )

    # ---- metricas por modelo (holdout, fold = -1) ----
    _sheet(
        wb, "metricas_por_modelo",
        ["modelo", "familia", "framework", "version", "activo", "rmse", "mae", "r2", "nse",
         "kge", "accuracy", "precision", "recall", "f1", "auc", "tn", "fp", "fn", "tp",
         "training_time_s"],
        [
            [
                r.name, r.algorithm_type, r.framework, r.version, r.is_active,
                r.rmse, r.mae, r.r2, r.nse, r.kge, r.accuracy, r.precision, r.recall, r.f1,
                r.auc, r.tn, r.fp, r.fn, r.tp, r.train_time_s,
            ]
            for r in rows_active
        ],
    )

    # ---- metricas por modelo y por fold (todas las filas de ml_model_metrics) ----
    fm = C.fold_metric_rows(db)
    _sheet(
        wb, "metricas_por_fold",
        ["modelo", "tarea", "version", "fold", "rmse", "mae", "r2", "nse", "kge", "accuracy",
         "precision", "recall", "f1", "auc_roc", "training_time_seconds"],
        [
            [
                m["modelo"], m["tarea"], m["version"], m["fold"], m["rmse"], m["mae"], m["r2"],
                m["nse"], m["kge"], m["accuracy"], m["precision"], m["recall"], m["f1"],
                m["auc_roc"], m["training_time_seconds"],
            ]
            for m in fm
        ],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
